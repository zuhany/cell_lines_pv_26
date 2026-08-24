import pandas as pd
import gseapy as gp
from openpyxl import load_workbook



# --------------------------------------------------
# Settings
# --------------------------------------------------

RANKED_FILE = "../data/processed/rnk.csv"
OUTPUT_DIR = "../results"

PERMUTATIONS = 10_000
SEED = 42


# --------------------------------------------------
# Helper functions
# --------------------------------------------------

# noinspection bad-argument-type
def run_gsea(ranked_genes, library_name, organism="Human"):
    """Run preranked GSEA yfor a given gene-set library."""

    gene_sets = gp.get_library(
        name=library_name,
        organism=organism,
    )

    print(f"\n{library_name}")
    print(f"Number of gene sets: {len(gene_sets)}")
    print(f"First 5 gene sets: {list(gene_sets.keys())[:5]}")

    gsea = gp.prerank(
        rnk=ranked_genes,
        gene_sets=gene_sets,
        #permutation_num=PERMUTATIONS,
        outdir=None,
        #seed=SEED,
        verbose=True,
        method='multilevel'
    )

    results = gsea.res2d[
        ["Term", "NES", "NOM p-val", "FDR q-val"]
    ].copy()

    return results.rename(
        columns={
            "Term": "PATHWAY",
            "NOM p-val": "NOMINAL_P",
            "FDR q-val": "FDR",
        }
    )


def save_results(results, output_file):
    """Save GSEA results to Excel with appropriate number formatting."""

    results.to_excel(output_file, index=False)

    wb = load_workbook(output_file)
    ws = wb.active

    if ws is None:
        raise RuntimeError(f"No active worksheet found in {output_file}")

    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "0.000000000"  # NES
        row[2].number_format = "0.000E+00"    # NOMINAL_P
        row[3].number_format = "0.000E+00"    # FDR

    wb.save(output_file)

    print(f"Saved: {output_file}")


# --------------------------------------------------
# Load ranked gene list
# --------------------------------------------------

rnk = pd.read_csv(
    RANKED_FILE,
    sep="\t",
    index_col=0,
)


# --------------------------------------------------
# Hallmark GSEA
# --------------------------------------------------

hallmark_results = run_gsea(
    ranked_genes=rnk,
    library_name="MSigDB_Hallmark_2020",
)

save_results(
    hallmark_results,
    f"{OUTPUT_DIR}/glass_GSEA_Hallmark_python.xlsx",
)


# --------------------------------------------------
# Reactome GSEA
# --------------------------------------------------

reactome_results = run_gsea(
    ranked_genes=rnk,
    library_name="Reactome_2022",
)

save_results(
    reactome_results,
    f"{OUTPUT_DIR}/glass_GSEA_Reactome_python.xlsx",
)